import os
from io import BytesIO
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class DocxExporter:
    """Exports structured question paper into an editable, standardized Microsoft Word (.docx) document."""

    @classmethod
    def generate_docx(cls, paper_data: Dict[str, Any], output_path: str = None) -> bytes:
        doc = Document()

        # Adjust margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        is_sgu = bool(
            (paper_data.get("institution_name") and "SANJAY GHODAWAT" in str(paper_data.get("institution_name")).upper()) or
            (paper_data.get("title") and "SGU" in str(paper_data.get("title")).upper()) or
            paper_data.get("layout") == "sgu"
        )

        # Header Titles
        if is_sgu:
            inst_name = paper_data.get("institution_name", "SANJAY GHODAWAT UNIVERSITY, KOLHAPUR")
            dept_name = paper_data.get("department", "School of Computer Science and Engineering")
            exam_title = paper_data.get("title", "END SEMESTER EXAMINATION")
        else:
            inst_name = paper_data.get("institution_name", "NATIONAL INSTITUTE OF TECHNOLOGY & SCIENCE")
            dept_name = paper_data.get("department", "Department of Computer Science & Engineering")
            exam_title = paper_data.get("title", "End Semester Theory Examination")

        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_inst = p_inst.add_run(inst_name.upper())
        run_inst.bold = True
        run_inst.font.size = Pt(13)
        run_inst.font.name = "Arial"

        p_dept = doc.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_dept = p_dept.add_run(dept_name)
        run_dept.bold = True
        run_dept.font.size = Pt(10)
        run_dept.font.name = "Arial"

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(exam_title)
        run_title.bold = True
        run_title.font.size = Pt(11)
        run_title.font.name = "Arial"

        # Metadata Table
        course_name = paper_data.get("course_name", "Computer Science" if is_sgu else "Data Structures & Algorithms")
        course_code = paper_data.get("course_code", "CSE601" if is_sgu else "CS301")
        semester = paper_data.get("semester", "VI" if is_sgu else "Semester V")
        duration = paper_data.get("duration_minutes", 180)
        total_marks = paper_data.get("total_marks", 100)
        academic_year = paper_data.get("academic_year", "2026–27")
        hours = duration // 60
        mins = duration % 60
        dur_str = f"{hours} Hours" + (f" {mins} Mins" if mins else "")

        meta_table = doc.add_table(rows=3, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.style = 'Table Grid'

        if is_sgu:
            labels = [
                ("Subject:", course_name, "Semester:", semester),
                ("Subject Code:", course_code, "Maximum Marks:", f"{total_marks} Marks"),
                ("Duration:", dur_str, "Academic Year:", academic_year)
            ]
        else:
            labels = [
                ("Course Code:", course_code, "Course Title:", course_name),
                ("Semester:", semester, "Max Marks:", f"{total_marks} Marks"),
                ("Duration:", dur_str, "Roll / Reg No:", "[                         ]")
            ]

        for r_idx, row in enumerate(labels):
            cells = meta_table.rows[r_idx].cells
            cells[0].paragraphs[0].add_run(row[0]).bold = True
            cells[1].paragraphs[0].add_run(row[1])
            cells[2].paragraphs[0].add_run(row[2]).bold = True
            cells[3].paragraphs[0].add_run(row[3])
            for c in cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.font.name = "Arial"

        # SGU Proposed Demo Notice
        if is_sgu:
            p_notice = doc.add_paragraph()
            p_notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_notice_b = p_notice.add_run("Notice: ")
            r_notice_b.bold = True
            r_notice_b.font.size = Pt(8.5)
            r_notice_b.font.color.rgb = RGBColor(180, 83, 9)
            r_notice = p_notice.add_run("This is a proposed demo format, not a verified official SGU examination pattern. Confirm the approved duration and marking scheme with your department.")
            r_notice.italic = True
            r_notice.font.size = Pt(8.5)
            r_notice.font.color.rgb = RGBColor(180, 83, 9)

        doc.add_paragraph() # Spacing

        # Instructions
        raw_instructions = paper_data.get("instructions") or (
            "1. Answer ALL questions according to instructions in respective sections.\n"
            "2. Marks for each question are indicated against it.\n"
            "3. Assume suitable data wherever necessary and state them clearly."
        )
        p_instr = doc.add_paragraph()
        p_instr.add_run("General Instructions:\n" if not is_sgu else "Instructions:\n").bold = True
        for line in raw_instructions.splitlines():
            if line.strip():
                p_instr.add_run(f"• {line.strip()}\n")
        p_instr.runs[0].font.size = Pt(9)

        # Sections
        show_parts = paper_data.get("show_part_headers", False) if "show_part_headers" in paper_data else (paper_data.get("paper_mode") == "sections")
        num_fmt = paper_data.get("numbering_format", "Q")

        for s_idx, sec in enumerate(paper_data.get("sections", [])):
            sec_title = sec.get("title", f"Q.{s_idx + 1} Answer the following")
            if is_sgu and not sec_title.upper().startswith("Q.") and not sec_title.upper().startswith("Q"):
                sec_title = f"Q.{s_idx + 1} {sec_title}"
            marks_per_q = sec.get("marks_per_question", 2)
            req_count = sec.get("required_count", len(sec.get("questions", [])))
            tot_sec_marks = sec.get("total_section_marks", req_count * marks_per_q)
            choice_type = sec.get("choice_type", "")

            if is_sgu:
                p_sec = doc.add_paragraph()
                r_sec = p_sec.add_run(f"{sec_title.upper()}")
                r_sec.bold = True
                r_sec.font.size = Pt(10.5)

                if choice_type and choice_type != "Compulsory":
                    r_sub = p_sec.add_run(f"  ({choice_type})")
                    r_sub.italic = True
                    r_sub.font.size = Pt(9.5)
                elif req_count and sec.get("questions") and req_count < len(sec.get("questions")):
                    r_sub = p_sec.add_run(f"  (Answer any {req_count} out of {len(sec.get('questions'))} questions × {marks_per_q} marks = {tot_sec_marks})")
                    r_sub.italic = True
                    r_sub.font.size = Pt(9.5)

                r_sec_marks = p_sec.add_run(f"\t[{tot_sec_marks}]")
                r_sec_marks.bold = True

                # Table for SGU subquestions
                q_table = doc.add_table(rows=0, cols=4)
                q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                q_table.style = 'Table Grid'

                for q_idx, q in enumerate(sec.get("questions", [])):
                    subq_letter = f"({chr(97 + q_idx)})"
                    row_cells = q_table.add_row().cells
                    row_cells[0].paragraphs[0].add_run(subq_letter).bold = True
                    
                    qp = row_cells[1].paragraphs[0]
                    qp.add_run(q.get("question_text", ""))
                    options = q.get("options", [])
                    if options and len(options) >= 2:
                        opt_letters = ['(a)', '(b)', '(c)', '(d)']
                        for o_idx, opt in enumerate(options[:4]):
                            qp.add_run(f"\n   {opt_letters[o_idx]} {opt}")
                    
                    meta_tag = f"[{q.get('blooms_level', 'Understand')}]"
                    row_cells[2].paragraphs[0].add_run(meta_tag)
                    row_cells[3].paragraphs[0].add_run(f"[{q.get('marks', marks_per_q)}]").bold = True

                    for c in row_cells:
                        for p in c.paragraphs:
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            for r in p.runs:
                                r.font.size = Pt(9)
                                r.font.name = "Arial"

                doc.add_paragraph()

            else:
                if show_parts:
                    p_sec = doc.add_paragraph()
                    r_sec = p_sec.add_run(f"{sec_title}  (Answer {req_count} × {marks_per_q} = {tot_sec_marks} Marks)")
                    r_sec.bold = True
                    r_sec.font.size = Pt(10.5)

                # Table for section questions
                q_table = doc.add_table(rows=1, cols=4)
                q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                q_table.style = 'Table Grid'
                hdr_cells = q_table.rows[0].cells
                hdr_cells[0].paragraphs[0].add_run("Q.No").bold = True
                hdr_cells[1].paragraphs[0].add_run("Question Description").bold = True
                hdr_cells[2].paragraphs[0].add_run("Bloom / CO").bold = True
                hdr_cells[3].paragraphs[0].add_run("Marks").bold = True

                for q in sec.get("questions", []):
                    raw_n = str(q.get("q_number", "")).lstrip("Qq. ")
                    if num_fmt == "num_only":
                        q_num = raw_n
                    elif num_fmt == "Q_dot":
                        q_num = f"Q.{raw_n}"
                    else:
                        q_num = f"Q{raw_n}"

                    row_cells = q_table.add_row().cells
                    row_cells[0].paragraphs[0].add_run(q_num).bold = True
                    
                    qp = row_cells[1].paragraphs[0]
                    qp.add_run(q.get("question_text", ""))
                    options = q.get("options", [])
                    if options and len(options) >= 2:
                        opt_letters = ['(a)', '(b)', '(c)', '(d)']
                        for o_idx, opt in enumerate(options[:4]):
                            qp.add_run(f"\n   {opt_letters[o_idx]} {opt}")
                    
                    meta_tag = f"[{q.get('blooms_level', 'L2')} / {q.get('difficulty', 'Med')}]"
                    row_cells[2].paragraphs[0].add_run(meta_tag)
                    row_cells[3].paragraphs[0].add_run(f"[{q.get('marks', marks_per_q)}]").bold = True

                    for c in row_cells:
                        for p in c.paragraphs:
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            for r in p.runs:
                                r.font.size = Pt(9)
                                r.font.name = "Arial"

                doc.add_paragraph()

        # Total Marks for SGU
        if is_sgu:
            p_sgu_tot = doc.add_paragraph()
            p_sgu_tot.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_sgu_tot = p_sgu_tot.add_run(f"TOTAL: {total_marks} MARKS")
            r_sgu_tot.bold = True
            r_sgu_tot.font.size = Pt(11)
            doc.add_paragraph()

        # End of paper
        p_end = doc.add_paragraph()
        p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_end = p_end.add_run("********** END OF QUESTION PAPER **********")
        r_end.bold = True
        r_end.font.size = Pt(9)
        r_end.font.color.rgb = RGBColor(100, 116, 139)

        if output_path:
            doc.save(output_path)
            with open(output_path, "rb") as f:
                return f.read()

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
