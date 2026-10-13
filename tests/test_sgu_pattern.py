import pytest
import io
from docx import Document
from backend.app.services.export_pdf import PDFExporter
from backend.app.services.export_docx import DocxExporter


def create_sgu_100_paper_data():
    return {
        "title": "END SEMESTER EXAMINATION",
        "institution_name": "SANJAY GHODAWAT UNIVERSITY, KOLHAPUR",
        "department": "School of Computer Science and Engineering",
        "semester": "Semester VI",
        "academic_year": "2026–2027",
        "course_code": "CSE601",
        "course_name": "Computer Science",
        "duration_minutes": 180,
        "total_marks": 100,
        "layout": "sgu",
        "sections": [
            {
                "section_id": "1",
                "title": "Q.1 Answer all questions (20 Marks)",
                "marks_per_question": 2,
                "required_count": 10,
                "total_section_marks": 20,
                "choice_type": "Compulsory",
                "questions": [
                    {"q_number": "1", "question_text": f"Explain basic computing concept #{i+1}.", "marks": 2, "blooms_level": "Remember"}
                    for i in range(10)
                ]
            },
            {
                "section_id": "2",
                "title": "Q.2 Answer the following (20 Marks)",
                "marks_per_question": 10,
                "required_count": 2,
                "total_section_marks": 20,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Analyze process scheduling algorithms in detail.", "marks": 10, "blooms_level": "Analyze"},
                    {"q_number": "2", "question_text": "Describe memory management paging and segmentation.", "marks": 10, "blooms_level": "Understand"},
                    {"q_number": "3", "question_text": "Evaluate deadlock prevention versus deadlock detection.", "marks": 10, "blooms_level": "Evaluate"}
                ]
            },
            {
                "section_id": "3",
                "title": "Q.3 Answer the following (20 Marks)",
                "marks_per_question": 10,
                "required_count": 2,
                "total_section_marks": 20,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Design a relational schema normalized up to 3NF and BCNF.", "marks": 10, "blooms_level": "Apply"},
                    {"q_number": "2", "question_text": "Compare ACID properties with BASE architecture.", "marks": 10, "blooms_level": "Analyze"},
                    {"q_number": "3", "question_text": "Explain indexing strategies using B+ trees.", "marks": 10, "blooms_level": "Understand"}
                ]
            },
            {
                "section_id": "4",
                "title": "Q.4 Answer the following (20 Marks)",
                "marks_per_question": 10,
                "required_count": 2,
                "total_section_marks": 20,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Implement Dijkstra's shortest path routing protocol.", "marks": 10, "blooms_level": "Apply"},
                    {"q_number": "2", "question_text": "Analyze TCP congestion control mechanisms: Tahoe vs Reno.", "marks": 10, "blooms_level": "Analyze"},
                    {"q_number": "3", "question_text": "Differentiate symmetric and asymmetric encryption systems.", "marks": 10, "blooms_level": "Understand"}
                ]
            },
            {
                "section_id": "5",
                "title": "Q.5 Answer the following (20 Marks)",
                "marks_per_question": 10,
                "required_count": 2,
                "total_section_marks": 20,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Formulate a convolutional neural network architecture for image classification.", "marks": 10, "blooms_level": "Create"},
                    {"q_number": "2", "question_text": "Explain transformer attention mechanisms and self-attention.", "marks": 10, "blooms_level": "Analyze"},
                    {"q_number": "3", "question_text": "Evaluate gradient descent optimization variants: Adam vs SGD.", "marks": 10, "blooms_level": "Evaluate"}
                ]
            }
        ]
    }


def create_sgu_50_paper_data():
    return {
        "title": "MID-TERM ASSESSMENT",
        "institution_name": "SANJAY GHODAWAT UNIVERSITY, KOLHAPUR",
        "department": "School of Computer Science and Engineering",
        "semester": "Semester VI",
        "academic_year": "2026–2027",
        "course_code": "CSE601",
        "course_name": "Computer Science",
        "duration_minutes": 120,
        "total_marks": 50,
        "layout": "sgu",
        "sections": [
            {
                "section_id": "1",
                "title": "Q.1 Answer all questions (10 Marks)",
                "marks_per_question": 2,
                "required_count": 5,
                "total_section_marks": 10,
                "choice_type": "Compulsory",
                "questions": [
                    {"q_number": "1", "question_text": f"Define term #{i+1}.", "marks": 2, "blooms_level": "Remember"}
                    for i in range(5)
                ]
            },
            *[
                {
                    "section_id": str(sec_i),
                    "title": f"Q.{sec_i} Answer the following (10 Marks)",
                    "marks_per_question": 5,
                    "required_count": 2,
                    "total_section_marks": 10,
                    "choice_type": "Answer any 2 out of 3 questions",
                    "questions": [
                        {"q_number": "1", "question_text": f"Explain topic {sec_i}A in detail.", "marks": 5, "blooms_level": "Understand"},
                        {"q_number": "2", "question_text": f"Compare topic {sec_i}B with alternatives.", "marks": 5, "blooms_level": "Apply"},
                        {"q_number": "3", "question_text": f"Analyze case scenario for topic {sec_i}C.", "marks": 5, "blooms_level": "Analyze"}
                    ]
                }
                for sec_i in range(2, 6)
            ]
        ]
    }


def create_sgu_25_paper_data():
    return {
        "title": "UNIT ASSESSMENT TEST",
        "institution_name": "SANJAY GHODAWAT UNIVERSITY, KOLHAPUR",
        "department": "School of Computer Science and Engineering",
        "semester": "Semester VI",
        "academic_year": "2026–2027",
        "course_code": "CSE601",
        "course_name": "Computer Science",
        "duration_minutes": 60,
        "total_marks": 25,
        "layout": "sgu",
        "sections": [
            {
                "section_id": "1",
                "title": "Q.1 Answer all questions (5 Marks)",
                "marks_per_question": 1,
                "required_count": 5,
                "total_section_marks": 5,
                "choice_type": "Compulsory",
                "questions": [
                    {"q_number": "1", "question_text": f"Quick test item #{i+1}.", "marks": 1, "blooms_level": "Remember"}
                    for i in range(5)
                ]
            },
            {
                "section_id": "2",
                "title": "Q.2 Answer the following (10 Marks)",
                "marks_per_question": 5,
                "required_count": 2,
                "total_section_marks": 10,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Explain Unit 1 core principles.", "marks": 5, "blooms_level": "Understand"},
                    {"q_number": "2", "question_text": "Illustrate Unit 1 practical implementation.", "marks": 5, "blooms_level": "Apply"},
                    {"q_number": "3", "question_text": "Evaluate Unit 1 performance tradeoffs.", "marks": 5, "blooms_level": "Analyze"}
                ]
            },
            {
                "section_id": "3",
                "title": "Q.3 Answer the following (10 Marks)",
                "marks_per_question": 5,
                "required_count": 2,
                "total_section_marks": 10,
                "choice_type": "Answer any 2 out of 3 questions",
                "questions": [
                    {"q_number": "1", "question_text": "Detail Unit 2 algorithmic procedure.", "marks": 5, "blooms_level": "Understand"},
                    {"q_number": "2", "question_text": "Demonstrate Unit 2 worked solution.", "marks": 5, "blooms_level": "Apply"},
                    {"q_number": "3", "question_text": "Critique Unit 2 boundary cases.", "marks": 5, "blooms_level": "Evaluate"}
                ]
            }
        ]
    }


def test_sgu_100_marks_balance():
    paper = create_sgu_100_paper_data()
    calculated_total = sum(sec["total_section_marks"] for sec in paper["sections"])
    assert calculated_total == 100
    assert paper["total_marks"] == 100
    assert len(paper["sections"]) == 5
    assert len(paper["sections"][0]["questions"]) == 10  # Q1 has 10 questions
    for sec in paper["sections"][1:]:
        assert len(sec["questions"]) == 3  # Q2-Q5 have 3 questions each
        assert sec["required_count"] == 2  # students answer any 2


def test_sgu_50_marks_balance():
    paper = create_sgu_50_paper_data()
    calculated_total = sum(sec["total_section_marks"] for sec in paper["sections"])
    assert calculated_total == 50
    assert paper["total_marks"] == 50
    assert len(paper["sections"]) == 5
    assert len(paper["sections"][0]["questions"]) == 5  # Q1 has 5 questions (5x2=10)
    for sec in paper["sections"][1:]:
        assert len(sec["questions"]) == 3  # Q2-Q5 have 3 questions each
        assert sec["required_count"] == 2  # 2 of 3 x 5m = 10m


def test_sgu_25_marks_balance():
    paper = create_sgu_25_paper_data()
    calculated_total = sum(sec["total_section_marks"] for sec in paper["sections"])
    assert calculated_total == 25
    assert paper["total_marks"] == 25
    assert len(paper["sections"]) == 3
    assert len(paper["sections"][0]["questions"]) == 5  # Q1 has 5 questions (5x1=5)
    for sec in paper["sections"][1:]:
        assert len(sec["questions"]) == 3  # Q2-Q3 have 3 questions each
        assert sec["required_count"] == 2  # 2 of 3 x 5m = 10m


def test_sgu_100_pdf_export():
    paper = create_sgu_100_paper_data()
    pdf_bytes = PDFExporter.generate_pdf(paper)
    assert pdf_bytes is not None
    assert len(pdf_bytes) > 1000
    # PDF starts with %PDF
    assert pdf_bytes.startswith(b"%PDF")


def test_sgu_50_pdf_export():
    paper = create_sgu_50_paper_data()
    pdf_bytes = PDFExporter.generate_pdf(paper)
    assert pdf_bytes is not None
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF")


def test_sgu_25_pdf_export():
    paper = create_sgu_25_paper_data()
    pdf_bytes = PDFExporter.generate_pdf(paper)
    assert pdf_bytes is not None
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF")


def test_sgu_100_docx_export():
    paper = create_sgu_100_paper_data()
    docx_bytes = DocxExporter.generate_docx(paper)
    assert docx_bytes is not None
    assert len(docx_bytes) > 1000
    # Read back docx to verify notice and headings
    doc = Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "SANJAY GHODAWAT UNIVERSITY" in all_text
    assert "Notice: This is a proposed demo format" in all_text
    assert "TOTAL: 100 MARKS" in all_text


def test_sgu_50_docx_export():
    paper = create_sgu_50_paper_data()
    docx_bytes = DocxExporter.generate_docx(paper)
    doc = Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "SANJAY GHODAWAT UNIVERSITY" in all_text
    assert "TOTAL: 50 MARKS" in all_text


def test_sgu_25_docx_export():
    paper = create_sgu_25_paper_data()
    docx_bytes = DocxExporter.generate_docx(paper)
    doc = Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "SANJAY GHODAWAT UNIVERSITY" in all_text
    assert "TOTAL: 25 MARKS" in all_text
